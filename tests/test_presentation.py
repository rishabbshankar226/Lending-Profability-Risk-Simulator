import csv
from dataclasses import replace
from decimal import Decimal as D
from io import BytesIO, StringIO
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook
import pytest

from lending_simulator.data import generate_dataset, load_dataset
from lending_simulator.decisions import compare_policies
from lending_simulator.presentation import cohort_heatmap, csv_package, workbook_payload, monthly_frame, assumption_register
from lending_simulator.exports import scenario_workbook
from lending_simulator.types import Assumptions
from lending_simulator.presentation import decision_brief


@pytest.fixture(scope="module")
def brief_dataset():
    return load_dataset(Path(__file__).resolve().parents[1] / "data" / "applications.csv")


@pytest.mark.parametrize("cash, loss, cash_margin, loss_margin", [
    ("50000", ".05", "0", "0"),
    ("49999", ".051", "-1", "-.1"),
    ("50001", ".049", "1", ".1"),
    ("50000.000000000000000000000000001", ".050000000000000000000000000000001",
     "1e-27", "-1e-31"),
    ("50000", None, "0", None),
])
def test_limit_margins_preserve_exact_boundaries_and_percentage_point_units(
        brief_dataset, cash, loss, cash_margin, loss_margin):
    from lending_simulator.presentation import limit_margins

    result = compare_policies(brief_dataset, Assumptions())[0]
    result = replace(result, summary=replace(
        result.summary, minimum_cash=D(cash), loss_ratio=None if loss is None else D(loss)))
    cash_gap, loss_gap = limit_margins(result)
    assert cash_gap == D(cash_margin)
    assert loss_gap == (None if loss_margin is None else D(loss_margin))


@pytest.mark.parametrize("assumptions, policy, recommendation, explanation", [
    (Assumptions(), "balanced", "Conservative", "highest expected full-runoff"),
    (replace(Assumptions(), initial_cash=D("1250000")), "balanced", "Balanced", "highest expected full-runoff"),
    (replace(Assumptions(), default_stress=D(2)), "conservative", "None", "No eligible strategy is profitable"),
    (replace(Assumptions(), initial_cash=D(0)), "conservative", "None", "No policy meets"),
])
def test_brief_reports_the_calculated_decision_and_exact_applied_inputs(
        brief_dataset, assumptions, policy, recommendation, explanation):
    results = compare_policies(brief_dataset, assumptions)
    selected = next(r for r in results if r.policy == policy)
    text = decision_brief(selected, results)
    assert f"**Recommended policy: {recommendation}.**" in text
    assert f"Selected portfolio: **{policy.title()}**" in text
    assert explanation in text
    assert "39 months" in text and "Dec 2029" in text
    assert "synthetic" in text and "uncalibrated" in text
    assert "positive operating profit" in text
    manifest = json.loads(text.split("```json\n", 1)[1].split("```", 1)[0])
    assert manifest["assumptions"] == assumptions.to_dict()
    assert manifest["run_id"] == selected.run_id
    assert manifest["dataset_hash"] == brief_dataset.dataset_hash
    assert manifest["comparison_run_ids"] == {r.policy: r.run_id for r in results}
    assert manifest["checks"]["passed"] is True
    if assumptions == Assumptions():
        assert ("| Aggressive | 14,989,287 | 396,446 | 2.29% | (1,485,855) | "
                "1,535,855 | Minimum-cash floor breached |") in text
        assert "Selected portfolio: **Balanced**" in text


@pytest.mark.parametrize("mismatch", ["selected", "comparison"])
def test_brief_rejects_stale_selected_results_and_mixed_comparison(brief_dataset, mismatch):
    base = compare_policies(brief_dataset, Assumptions())
    capital = compare_policies(brief_dataset, replace(Assumptions(), initial_cash=D("1250000")))
    selected, results = (base[1], capital) if mismatch == "selected" else (base[0], (base[0], capital[1], base[2]))
    with pytest.raises(ValueError):
        decision_brief(selected, results)


def test_csv_package_has_manifest_units_exact_values_and_matching_run():
    dataset = generate_dataset(count=100)
    result = compare_policies(dataset, Assumptions())[1]
    archive = zipfile.ZipFile(BytesIO(csv_package(result, dataset)))
    manifest = json.loads(archive.read("manifest.json"))
    assert manifest["run_id"] == result.run_id
    assert manifest["dataset_hash"] == dataset.dataset_hash
    assert manifest["assumptions"] == result.assumptions.to_dict()
    assert manifest["currency"] == "USD"
    rows = list(csv.DictReader(StringIO(archive.read("monthly.csv").decode())))
    assert len(rows) == 39
    assert D(rows[0]["ending_cash"]) == result.monthly[0].ending_cash
    assert json.loads(archive.read("units.json"))["monthly"]["ending_cash"] == "USD, month-end balance"
    assert json.loads(archive.read("units.json"))["summary"]["operating_result"] == "USD, full-runoff total"


def test_heatmap_preserves_future_unobserved_ages_as_missing():
    result = compare_policies(generate_dataset(count=100), Assumptions())[1]
    heatmap = cohort_heatmap(result, as_of_month=23)
    assert heatmap.loc[23, 0] == 0
    assert heatmap.loc[23, 1:].isna().all()
    assert heatmap.loc[0, 12] > 0
    assert monthly_frame(result)["ending_cash"].iloc[0] == float(result.monthly[0].ending_cash)


def test_workbook_contains_selected_scenario_and_independent_formula_benchmarks():
    dataset = generate_dataset(count=100)
    assumptions = replace(Assumptions(), funding_rate=D(".12"))
    results = compare_policies(dataset, assumptions)
    result = results[1]
    book = load_workbook(BytesIO(scenario_workbook(result, results, dataset)), data_only=False)
    cached = load_workbook(BytesIO(scenario_workbook(result, results, dataset)), data_only=True)
    assert book["Assumptions"]["B2"].value == result.run_id
    assert cached["Summary"]["B5"].value == float(result.summary.operating_result)
    assert book["Loan benchmark"]["B8"].value.startswith("=IF(")
    assert book["Loan benchmark"]["B6"].number_format.endswith("%")
    assert book["Loan benchmark"]["F23"].value.startswith("=")
    assert abs(cached["Loan benchmark"]["F23"].value) < .01
    assert cached["Default benchmark"]["B7"].value == 675
    assert cached["Default benchmark"]["E16"].value == 225
    assert cached["Default benchmark"]["E15"].value == 0
    assert workbook_payload(result, results, dataset)["selected_run_id"] == result.run_id


def test_changed_assumptions_export_a_new_scenario_without_mutating_original():
    dataset = generate_dataset(count=100)
    base = compare_policies(dataset, Assumptions())[1]
    changed = compare_policies(dataset, replace(Assumptions(), default_stress=D(2)))[1]
    assert base.run_id != changed.run_id
    base_archive = zipfile.ZipFile(BytesIO(csv_package(base, dataset)))
    changed_archive = zipfile.ZipFile(BytesIO(csv_package(changed, dataset)))
    assert json.loads(base_archive.read("manifest.json"))["assumptions"]["default_stress"] == "1"
    assert json.loads(changed_archive.read("manifest.json"))["assumptions"]["default_stress"] == "2"


def test_assumptions_table_serializes_to_the_dashboard_without_arrow_type_fixes():
    import pyarrow as pa
    import pandas as pd
    table = pa.Table.from_pandas(pd.DataFrame(assumption_register(Assumptions())))
    assert table.num_rows == 19
