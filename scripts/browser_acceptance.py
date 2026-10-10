"""Exercise the real candidate in isolated Chromium; retain reviewable evidence.

Run with requirements-browser.txt and `playwright install chromium` installed.
The server binds only to loopback and runs app.py unchanged. No hosting account,
production deployment, or application-state injection is used. The empty-input
case uses a separate valid zero-row dataset with unchanged application/model code.
"""

import argparse
from contextlib import contextmanager
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import traceback
from urllib.request import urlopen
import zipfile

from playwright.sync_api import expect, sync_playwright


ROOT = Path(__file__).resolve().parents[1]
VIEWS = ("Overview", "Policies", "Cohorts", "Funding & stress", "Methodology")
BASE_RUN = "5a88f7bef3fa93d4"
BALANCED_RUN = "a2feec5a72facbc3"


class Review:
    def __init__(self, page, output, name, report):
        self.page, self.output, self.name, self.report = page, output, name, report
        self.console = []
        self.requests = []
        self.websockets = []
        page.on("pageerror", lambda e: self.console.append({"type": "pageerror", "text": str(e)}))
        page.on("console", lambda m: self.console.append({"type": m.type, "text": m.text})
                if m.type in ("error", "warning") else None)
        page.on("requestfailed", lambda r: self.requests.append({"url": r.url, "failure": r.failure}))
        page.on("websocket", lambda w: self.websockets.append(w.url))

    def check(self, name, details=None):
        self.report["checks"].append({"name": name, "details": details})
        print(f"PASS {self.name}: {name}", flush=True)

    def ready(self):
        expect(self.page.get_by_role("heading", name="Lending workbench", exact=True)).to_be_visible()
        expect(self.page.get_by_test_id("stMetricValue").first).to_be_visible()
        self.page.wait_for_function("""() => !Array.from(document.querySelectorAll('[data-testid="stStatusWidget"]'))
            .some(e => e.getClientRects().length && /running|connecting/i.test(e.innerText))""")
        self.page.wait_for_function("() => !document.querySelector('[data-stale=\"true\"]')")
        if self.page.get_by_test_id("stPlotlyChart").count():
            self.page.wait_for_function("""() => Array.from(document.querySelectorAll('[data-testid="stPlotlyChart"]'))
                .every(e => e.querySelector('.js-plotly-plot')?.data?.length)""")
        assert self.page.get_by_test_id("stException").count() == 0

    def metrics(self):
        return self.page.get_by_test_id("stMetricValue").all_inner_texts()[:3]

    def capture(self, suffix, target=None):
        self.ready()
        if target is not None:
            target.scroll_into_view_if_needed()
        self.page.screenshot(path=str(self.output / f"{self.name}-{suffix}.png"), full_page=True)
        self.report["captures"].append(f"{self.name}-{suffix}.png")

    def navigate(self, view):
        start = time.monotonic()
        self.page.get_by_role("radio", name=view, exact=True).click()
        expect(self.page.get_by_role("heading", name=view, exact=True)).to_be_visible()
        self.ready()
        self.report["operations"].append({"name": "Navigate to " + view, "seconds": time.monotonic() - start})

    def expand(self, title):
        self.page.get_by_text(title, exact=True).click()

    def policy(self, name):
        self.select("Viewed policy", name)
        expect(self.page.get_by_text(f"Viewing policy: {name}", exact=True)).to_be_visible()
        self.ready()

    def select(self, label, name):
        self.ready()
        combo = self.page.get_by_role("combobox", name=label, exact=True)
        combo.scroll_into_view_if_needed()
        combo.click()
        expect(combo).to_have_attribute("aria-expanded", "true")
        self.page.get_by_role("option", name=name, exact=True).click()
        self.ready()

    def open_sidebar(self):
        sidebar = self.page.get_by_test_id("stSidebar")
        if sidebar.get_attribute("aria-expanded") != "true":
            self.page.get_by_test_id("stExpandSidebarButton").click()
        expect(sidebar).to_have_attribute("aria-expanded", "true")
        expect(self.page.get_by_role("heading", name="Scenario", exact=True)).to_be_visible()

    def close_sidebar(self):
        self.page.get_by_test_id("stSidebarCollapseButton").get_by_role("button").click()
        expect(self.page.get_by_test_id("stSidebar")).to_have_attribute("aria-expanded", "false")

    def reset(self):
        self.open_sidebar()
        self.page.get_by_role("button", name="Reset to base", exact=True).click()
        expect(self.page.get_by_test_id("stMetricValue").first).to_have_text("$48,133")
        self.ready()

    def applied_run(self):
        self.page.get_by_role("button", name="Applied inputs", exact=True).click()
        text = self.page.get_by_text(re.compile(r"Applied run: [0-9a-f]{16}")).inner_text()
        run = re.search(r"Applied run: ([0-9a-f]{16})", text).group(1)
        self.page.keyboard.press("Escape")
        return run

    def geometry(self):
        result = self.page.evaluate("""() => {
            const main = document.querySelector('[data-testid="stMain"]') || document.documentElement;
            const r = main.getBoundingClientRect();
            const painted = e => {
                for(let node=e;node;node=node.parentElement){
                    const style=getComputedStyle(node);
                    if(style.visibility!=='visible'||Number(style.opacity)===0)return false;
                    if(node.tagName==='DETAILS'&&!node.open&&!node.querySelector('summary')?.contains(e))return false;
                }
                return true;
            };
            const overflow = Array.from(main.querySelectorAll('h1,h2,h3,button,.workbench-status,.policy-mobile article'))
              .filter(e => e.getClientRects().length && painted(e) && getComputedStyle(e).position !== 'fixed')
              .map(e => ({e, r:e.getBoundingClientRect()}))
              .filter(x => x.r.width && (x.r.left < r.left-2 || x.r.right > r.right+2))
              .map(x => ({tag:x.e.tagName,text:x.e.innerText.slice(0,120),left:x.r.left,right:x.r.right}));
            return {viewport:innerWidth, mainWidth:r.width, scrollWidth:main.scrollWidth,
                clientWidth:main.clientWidth, overflow,
                tableVisible:Array.from(document.querySelectorAll('.policy-table')).some(e=>e.getClientRects().length),
                cardsVisible:Array.from(document.querySelectorAll('.policy-mobile')).some(e=>e.getClientRects().length)};
        }""")
        assert result["scrollWidth"] <= result["clientWidth"] + 2, result
        assert not result["overflow"], result
        return result

    def charts(self):
        return self.page.locator(".js-plotly-plot").evaluate_all("""plots => plots.map(p => ({
            width:p.getBoundingClientRect().width, height:p.getBoundingClientRect().height,
            traces:(p.data||[]).map(t=>({type:t.type,name:t.name,points:t.x?.length,
              first:t.x?.[0],last:t.x?.[t.x.length-1],zmin:t.zmin,zmax:t.zmax,
              dash:t.line?.dash,marker:t.marker?.symbol,hoverongaps:t.hoverongaps})),
            shapes:p.layout?.shapes, xaxis:p.layout?.xaxis, yaxis:p.layout?.yaxis,
            font:p.layout?.font, legend:p.layout?.legend
        }))""")

    def caption_contrast(self):
        samples = self.page.locator(".workbench-caption").evaluate_all("""els => els.filter(e=>e.getClientRects().length).map(e=>{
            let opacity=1, bg='rgb(245, 247, 250)';
            for(let n=e;n;n=n.parentElement){
                const s=getComputedStyle(n); opacity*=Number(s.opacity);
            }
            for(let n=e;n;n=n.parentElement){
                const c=getComputedStyle(n).backgroundColor;
                if(c!=='rgba(0, 0, 0, 0)'&&c!=='transparent'){bg=c;break;}
            }
            return {text:e.innerText.slice(0,90),foreground:getComputedStyle(e).color,background:bg,opacity};
        })""")
        assert samples, "No supporting text rendered"
        def rgb(value):
            return [float(x) for x in re.findall(r"[0-9.]+", value)][:3]
        def luminance(channels):
            c = [x/255 for x in channels]
            c = [x/12.92 if x <= .04045 else ((x+.055)/1.055)**2.4 for x in c]
            return sum(x*y for x,y in zip(c,(.2126,.7152,.0722)))
        for sample in samples:
            foreground, background = rgb(sample["foreground"]), rgb(sample["background"])
            blended = [f*sample["opacity"] + b*(1-sample["opacity"]) for f,b in zip(foreground,background)]
            values = sorted((luminance(blended),luminance(background)))
            sample["ratio"] = (values[1]+.05)/(values[0]+.05)
            assert sample["ratio"] >= 4.5, sample
        return {"minimum_ratio": min(s["ratio"] for s in samples), "samples": samples}


def desktop_journey(r):
    p = r.page
    assert r.metrics() == ["$48,133", "$97,450", "4.16 pp"]
    assert r.applied_run() == BASE_RUN
    r.capture("overview")
    chart = r.charts()
    assert chart[0]["traces"][0]["points"] == 39, chart
    assert chart[0]["xaxis"]["type"] == "date"
    assert any(s.get("y0") == 50000 for s in chart[1]["shapes"]), chart
    r.check("Base identity, 39 monthly points, calendar axis and cash floor", chart)
    r.capture("overview-charts", p.locator(".js-plotly-plot").first)
    r.expand("How operating profit is calculated")
    expect(p.get_by_role("columnheader", name="Component", exact=True)).to_be_visible()
    r.capture("profit-bridge", p.get_by_role("columnheader", name="Component", exact=True))
    r.capture("profit-bridge-chart", p.locator(".js-plotly-plot").nth(2))
    r.check("Profit bridge and numeric alternative open")

    equity = p.get_by_role("spinbutton", name="Starting equity cash ($)", exact=True)
    equity.fill("1250000")
    # No Enter/blur helper: this is the real type-then-click-Run race.
    p.get_by_role("button", name="Run scenario", exact=True).click()
    expect(p.get_by_role("heading", name="Recommended policy: Balanced", exact=True)).to_be_visible()
    expect(p.get_by_text("Unapplied changes · Results use the last successful scenario.", exact=True)).not_to_be_visible()
    r.check("Typed numeric edit commits when Run is clicked immediately")
    r.policy("Balanced")
    assert r.applied_run() == "690e9fa563e2d475"
    assert r.metrics()[0:2] == ["$294,468", "$3,857"]
    r.capture("capital")
    r.check("More-capital recommendation and viewed run remain distinct")

    equity.fill("1500000")
    equity.press("Tab")
    expect(p.get_by_text("Unapplied changes · Results use the last successful scenario.", exact=True)).to_be_visible()
    before = r.metrics()
    r.navigate("Policies")
    assert r.metrics() == before
    p.get_by_role("button", name="Restore applied inputs", exact=True).click()
    expect(equity).to_have_value("1250000.00")
    expect(p.get_by_text("Unapplied changes · Results use the last successful scenario.", exact=True)).not_to_be_visible()
    r.check("Draft survives navigation; Restore retains applied run")
    r.capture("policies")
    r.capture("policy-table", p.locator(".policy-table"))
    r.check("Desktop policy table geometry", r.geometry())

    r.expand("Compare with a pinned scenario")
    p.get_by_role("button", name="Pin applied scenario", exact=True).click()
    expect(p.get_by_text(re.compile(r"Baseline runs: .*Balanced 690e9fa563e2d475"))).to_be_visible()
    r.ready()
    r.expand("Credit assumptions")
    lag = p.get_by_role("spinbutton", name="Recovery lag (months)", exact=True)
    lag.fill("12")
    p.get_by_role("button", name="Run scenario", exact=True).click()
    expect(p.get_by_text(re.compile(r"Runoff differs: baseline 39 months, current 48 months"))).to_be_visible()
    r.capture("different-runoff")
    r.capture("runoff-comparison", p.get_by_text(re.compile(r"Runoff differs: baseline 39 months, current 48 months")))
    r.check("Pinned comparison distinguishes 39- and 48-month runoff")

    r.reset()
    r.navigate("Cohorts")
    cohort = r.charts()
    p.get_by_text("Complete runoff", exact=True).click()
    expect(p.get_by_role("radio", name="Complete runoff", exact=True)).to_be_checked()
    r.ready()
    assert r.charts()[0]["traces"][0]["zmax"] == cohort[0]["traces"][0]["zmax"]
    r.capture("cohorts")
    r.capture("cohort-chart", p.locator(".js-plotly-plot").first)
    r.check("Cohort cutoff retains shared scale and blank-gap hover behavior", r.charts())

    r.navigate("Funding & stress")
    p.get_by_role("button", name="Run stress grid", exact=True).click()
    expect(p.get_by_role("combobox", name="Inspect a stress case", exact=True)).to_be_visible()
    r.select("Inspect a stress case", "0.5× defaults · 12.00% annual funding")
    r.ready()
    stress = r.charts()[1]
    assert stress["xaxis"]["type"] == stress["yaxis"]["type"] == "category", stress
    assert stress["traces"][0]["first"] == "4.00%" and stress["traces"][0]["last"] == "16.00%", stress
    r.check("Stress axes retain percentage and multiplier categories", stress)
    identity = r.applied_run()
    p.get_by_role("button", name="Stage case as draft", exact=True).click()
    expect(p.get_by_text("Unapplied changes · Results use the last successful scenario.", exact=True)).to_be_visible()
    assert r.applied_run() == identity
    expect(p.get_by_role("button", name="Stage case as draft", exact=True)).to_be_disabled()
    r.capture("staged-stress")
    r.capture("stress-chart", p.locator(".js-plotly-plot").nth(1))
    r.check("Stress staging changes only draft and guards an existing draft")
    p.get_by_role("button", name="Run scenario", exact=True).click()
    expect(p.get_by_text("Stress grid has not been run for this applied policy. Run it when you want to explore the cases.", exact=True)).to_be_visible()
    r.check("Applying staged stress compares policies and clears stale grid")

    r.reset()
    r.expand("One-click examples")
    p.get_by_role("button", name="Higher defaults · 2×", exact=True).click()
    expect(p.get_by_role("heading", name="No positive-profit recommendation", exact=True)).to_be_visible()
    assert r.applied_run() == "a2d0056e84a3ceb8"
    r.capture("higher-defaults")
    r.check("Loss-making example has no positive-profit recommendation")

    r.reset()
    equity.fill("-1")
    p.get_by_role("button", name="Run scenario", exact=True).click()
    expect(p.get_by_text(re.compile(r"Scenario could not run:"))).to_be_visible()
    assert r.metrics()[0] == "$48,133"
    assert r.applied_run() == BASE_RUN
    r.capture("invalid-input", p.get_by_text(re.compile(r"Scenario could not run:")))
    r.check("Invalid inputs retain last successful scenario")
    p.get_by_role("button", name="Restore applied inputs", exact=True).click()

    r.navigate("Methodology")
    expect(p.get_by_text("Financial reconciliations passed", exact=True)).to_be_visible()
    r.capture("methodology")
    r.check("Methodology exposes reconciliation status without runtime error")


def downloads(r):
    p = r.page
    p.get_by_role("button", name="Export", exact=True).click()
    expect(p.get_by_role("button", name="Decision brief", exact=True)).to_be_visible()
    r.check("Export opens without triggering a browser download")
    for label, suffix in (("Decision brief", ".md"), ("Audit workbook", ".xlsx"), ("CSV results + manifest", ".zip")):
        with p.expect_download() as request:
            p.get_by_role("button", name=label, exact=True).click()
        download = request.value
        assert download.failure() is None
        assert BASE_RUN in download.suggested_filename
        assert download.suggested_filename.endswith(suffix)
        path = r.output / download.suggested_filename
        download.save_as(path)
        if suffix == ".md":
            assert BASE_RUN in path.read_text()
        elif suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                manifest = json.loads(archive.read("manifest.json"))
                assert manifest["run_id"] == BASE_RUN
        else:
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            assert "Assumptions" in workbook.sheetnames
            assert workbook["Assumptions"]["B2"].value == BASE_RUN
            workbook.close()
        r.check(f"Actual {label} transport and captured run identity", {"file": path.name, "bytes": path.stat().st_size})
    r.expand("Preview decision brief")
    expect(p.get_by_role("heading", name="Lending decision brief", exact=True)).to_be_visible()
    r.capture("export-preview")
    p.keyboard.press("Escape")
    r.policy("Balanced")
    assert r.applied_run() == BALANCED_RUN
    p.get_by_role("button", name="Export", exact=True).click()
    with p.expect_download() as request:
        p.get_by_role("button", name="Decision brief", exact=True).click()
    download = request.value
    path = r.output / download.suggested_filename
    download.save_as(path)
    assert BALANCED_RUN in path.name and BALANCED_RUN in path.read_text()
    r.check("Policy switch requests a newly identified export")


def responsive(r):
    p = r.page
    mobile = p.viewport_size["width"] <= 640
    r.capture("overview", p.get_by_role("heading", name="Lending workbench", exact=True))
    r.check("First-screen reflow", r.geometry())
    r.check("Supporting text rendered with at least 4.5:1 contrast", r.caption_contrast())
    nav = p.get_by_role("radio", name="Policies", exact=True).bounding_box()
    assert nav["y"] + nav["height"] <= p.viewport_size["height"], nav
    r.check("Analysis navigation is visible in the initial viewport", nav)
    for view in VIEWS[1:]:
        r.navigate(view)
        details = r.geometry()
        if view == "Policies":
            assert details["cardsVisible"] == mobile, details
            assert details["tableVisible"] != mobile, details
            r.capture("policy-comparison", p.locator(".policy-mobile" if mobile else ".policy-table"))
        if view != "Methodology":
            r.capture("chart-" + view.lower().replace(" & ", "-"), p.locator(".js-plotly-plot").first)
        r.capture(view.lower().replace(" & ", "-").replace(" ", "-"),
                  p.get_by_role("heading", name=view, exact=True))
        r.check(f"{view} reflow and visible controls", details)
    if mobile:
        r.open_sidebar()
        expect(p.get_by_role("spinbutton", name="Starting equity cash ($)", exact=True)).to_be_visible()
        p.get_by_role("button", name="Run scenario", exact=True).scroll_into_view_if_needed()
        r.capture("sidebar-run")
        r.close_sidebar()
        r.navigate("Policies")
        p.get_by_role("button", name="Export", exact=True).click()
        expect(p.get_by_role("button", name="Audit workbook", exact=True)).to_be_visible()
        r.capture("reachable-export")
        p.keyboard.press("Escape")
        r.check("Phone sidebar opens, reaches Run, returns to analysis and opens exports")


def download_switch(r):
    p = r.page
    p.get_by_role("button", name="Export", exact=True).click()
    with p.expect_download() as event:
        p.get_by_role("button", name="CSV results + manifest", exact=True).click()
    download = event.value
    p.keyboard.press("Escape")
    r.policy("Balanced")
    assert r.applied_run() == BALANCED_RUN
    assert download.failure() is None
    path = r.output / download.suggested_filename
    download.save_as(path)
    assert BASE_RUN in path.name
    with zipfile.ZipFile(path) as package:
        assert json.loads(package.read("manifest.json"))["run_id"] == BASE_RUN
    r.check("Requested CSV retains original filename and manifest across a policy switch", {
        "requested_run": BASE_RUN, "now_viewed_run": BALANCED_RUN, "file": path.name,
        "scope": "Actual download requested before switch, saved after switch. Loopback transfer may complete before selection; generation capture has separate unit coverage."})
    r.capture("completed")


def empty_portfolio(r):
    p = r.page
    expect(p.get_by_role("heading", name="No eligible policy", exact=True)).to_be_visible()
    expect(p.get_by_test_id("stMetricValue").nth(2)).to_have_text("Unavailable")
    r.navigate("Cohorts")
    expect(p.get_by_text("No funded cohorts under this policy. Cohort loss ratios are unavailable.", exact=True)).to_be_visible()
    r.capture("cohorts", p.get_by_role("heading", name="Cohorts", exact=True))
    r.check("Empty portfolio distinguishes unavailable ratios from zero")
    r.navigate("Funding & stress")
    p.get_by_role("button", name="Run stress grid", exact=True).click()
    expect(p.get_by_role("combobox", name="Stress result", exact=True)).to_be_visible()
    r.select("Stress result", "Net principal loss")
    expect(p.get_by_text("Net loss is unavailable: no principal was funded. Costs, cash, and failure reasons remain available.", exact=True)).to_be_visible()
    r.capture("stress", p.get_by_role("heading", name="Default and funding stress", exact=True))
    r.check("Empty-portfolio stress inspection renders unavailable net loss without exception")


def download_retry(r):
    p = r.page
    session = p.context.new_cdp_session(p)
    context_id = session.send("Target.getTargetInfo")["targetInfo"]["browserContextId"]
    session.send("Browser.setDownloadBehavior", {"behavior": "deny", "browserContextId": context_id, "eventsEnabled": True})
    p.get_by_role("button", name="Export", exact=True).click()
    with p.expect_download() as event:
        p.get_by_role("button", name="CSV results + manifest", exact=True).click()
    failure = event.value.failure()
    assert failure is not None, "The deliberately denied download unexpectedly succeeded"
    session.send("Browser.setDownloadBehavior", {"behavior": "allowAndName", "browserContextId": context_id,
        "downloadPath": str(r.output / "native-downloads"), "eventsEnabled": True})
    p.keyboard.press("Escape")
    assert r.applied_run() == BASE_RUN
    assert r.metrics() == ["$48,133", "$97,450", "4.16 pp"]
    p.get_by_role("button", name="Export", exact=True).click()
    with p.expect_download() as event:
        p.get_by_role("button", name="CSV results + manifest", exact=True).click()
    download = event.value
    assert download.failure() is None
    path = r.output / ("retried-" + download.suggested_filename)
    download.save_as(path)
    with zipfile.ZipFile(path) as package:
        assert json.loads(package.read("manifest.json"))["run_id"] == BASE_RUN
    media_url = re.compile(r"^http://127\.0\.0\.1:\d+/media/")
    assert all(media_url.match(item["url"]) for item in r.requests), r.requests
    r.report["intentional_failed_requests"] = r.requests.copy()
    r.requests.clear()
    r.check("Canceled download preserves scenario and retries with the same captured identity", {
        "failure": failure, "retried_file": path.name, "run_id": BASE_RUN,
        "scope": "Chromium temporarily denies this isolated context's download, then permits retry; no financial output is mocked. Server generation failures are outside this check."})
    session.detach()


@contextmanager
def live_server(root, log_path, *, entrypoint=None):
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    address = f"http://127.0.0.1:{port}"
    with log_path.open("w") as log:
        server = subprocess.Popen([sys.executable, "-m", "streamlit", "run", str(entrypoint or root / "app.py"),
            "--server.headless=true", "--server.address=127.0.0.1", f"--server.port={port}",
            "--browser.gatherUsageStats=false"], cwd=root, stdout=log, stderr=subprocess.STDOUT)
        try:
            deadline = time.monotonic() + 45
            while True:
                try:
                    with urlopen(address + "/_stcore/health", timeout=1) as health:
                        if health.status == 200:
                            break
                except OSError:
                    pass
                if server.poll() is not None or time.monotonic() > deadline:
                    raise RuntimeError(f"Streamlit did not start; inspect {log_path.name}")
                time.sleep(.2)
            yield address
        finally:
            server.terminate()
            try:
                server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                server.kill()
                server.wait(timeout=10)


def empty_dataset_source(destination):
    """Exercise real validation/rendering on a separate valid, empty input source."""
    from lending_simulator.data import Dataset, save_dataset
    shutil.copy2(ROOT / "app.py", destination / "app.py")
    for directory in ("lending_simulator", ".streamlit"):
        shutil.copytree(ROOT / directory, destination / directory,
                        ignore=shutil.ignore_patterns("__pycache__"))
    dataset = Dataset((), 2262026)
    target = destination / "data" / "applications.csv"
    target.parent.mkdir()
    save_dataset(dataset, target)
    (target.parent / "manifest.json").write_text(json.dumps(dataset.manifest()))
    return dataset.manifest()


def keyboard(r):
    p = r.page
    order = []
    for _ in range(45):
        p.keyboard.press("Tab")
        item = p.evaluate("""() => {const e=document.activeElement;return {tag:e.tagName,role:e.getAttribute('role'),
            label:e.getAttribute('aria-label')||e.innerText?.slice(0,100)||'',type:e.getAttribute('type')};}""")
        order.append(item)
        if item["role"] == "radio" and item["label"] == "Overview":
            break
    assert any(item["label"] == "Run scenario" for item in order), order
    assert any(item["label"] == "Viewed policy" for item in order), order
    r.check("Normal Tab order reaches scenario Run, policy selector and analysis", order)
    p.get_by_role("radio", name="Policies", exact=True).focus()
    p.keyboard.press("Enter")
    expect(p.get_by_role("heading", name="Policies", exact=True)).to_be_visible()
    assert p.get_by_role("radio", name="Policies", exact=True).evaluate("e=>e===document.activeElement")
    r.check("Keyboard navigation activates selected view and preserves focus")
    p.get_by_role("button", name="Applied inputs", exact=True).focus()
    p.keyboard.press("Enter")
    expect(p.get_by_text(re.compile("Applied run: " + BASE_RUN))).to_be_visible()
    p.keyboard.press("Escape")
    assert p.get_by_role("button", name="Applied inputs", exact=True).evaluate("e=>e===document.activeElement")
    r.check("Applied-input popover opens and returns focus using keyboard")
    focus = p.get_by_role("button", name="Applied inputs", exact=True).evaluate("""e=>{
        const c=getComputedStyle(e); return {outline:c.outline,boxShadow:c.boxShadow,
          focused:e===document.activeElement,visible:!!e.getClientRects().length};}""")
    assert focus["focused"] and focus["visible"]
    assert focus["outline"] not in ("none", "rgb(0, 0, 0) none 0px") or focus["boxShadow"] != "none", focus
    r.capture("keyboard-focus")
    r.check("Visible focused control and reduced-motion media", {"focus": focus,
        "reduced_motion": p.evaluate("matchMedia('(prefers-reduced-motion: reduce)').matches")})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts" / "browser-acceptance")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    source_files = [ROOT / "app.py", ROOT / ".streamlit" / "config.toml", *sorted((ROOT / "lending_simulator" / "ui").glob("*.py"))]
    report = {"created_at": datetime.now(timezone.utc).isoformat(), "source_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "source_sha256": {str(p.relative_to(ROOT)): sha256(p.read_bytes()).hexdigest() for p in source_files},
        "scope": "Isolated headless Chromium against unchanged app.py on loopback; synthetic dataset. Mobile sizes are emulation, not physical-device or intended-host measurements.",
        "cases": []}
    with tempfile.TemporaryDirectory(prefix="lending-browser-empty-") as directory:
        empty_root = Path(directory)
        report["empty_dataset_case"] = {"source": "Separate copy of unchanged app/code/config with a valid zero-row Dataset and manifest.",
            "manifest": empty_dataset_source(empty_root)}
        (output / "report.json").write_text(json.dumps(report, indent=2))
        with live_server(ROOT, output / "streamlit.log") as address, \
             live_server(empty_root, output / "streamlit-empty.log") as empty_address, \
             sync_playwright() as engine:
            browser = engine.chromium.launch(downloads_path=str(output / "native-downloads"))
            report["browser"] = browser.version
            for name, size, task, options in (
                ("desktop-1440", (1440, 900), desktop_journey, {}),
                ("desktop-1280", (1280, 800), responsive, {}),
                ("downloads", (1440, 900), downloads, {}),
                ("phone-390", (390, 844), responsive, {"is_mobile": True, "has_touch": True}),
                ("reflow-320", (320, 800), responsive, {"is_mobile": True, "has_touch": True}),
                ("keyboard", (1280, 800), keyboard, {"reduced_motion": "reduce"}),
                ("download-switch", (1440, 900), download_switch, {}),
                ("empty-portfolio", (1440, 900), empty_portfolio, {}),
                ("download-retry", (1440, 900), download_retry, {}),
            ):
                case = {"name": name, "viewport": size, "checks": [], "captures": [], "operations": []}
                report["cases"].append(case)
                context = browser.new_context(viewport={"width": size[0], "height": size[1]}, accept_downloads=True, **options)
                context.tracing.start(screenshots=True, snapshots=True, sources=True)
                page = context.new_page()
                page.set_default_timeout(15000)
                r = Review(page, output, name, case)
                try:
                    start = time.monotonic()
                    page.goto(empty_address if name == "empty-portfolio" else address, wait_until="domcontentloaded")
                    r.ready()
                    case["operations"].append({"name": "Initial candidate render", "seconds": time.monotonic() - start})
                    task(r)
                    errors = [e for e in r.console if e["type"] in ("error", "pageerror")]
                    assert not errors, errors
                    assert not r.requests, r.requests
                    assert r.websockets, "No Streamlit WebSocket observed"
                    r.check("No unexpected browser errors or failed requests; Streamlit WebSocket connected")
                    case["status"] = "passed"
                except Exception:
                    case["status"] = "failed"
                    case["error"] = traceback.format_exc()
                    print(f"FAIL {name}: {case['error']}", flush=True)
                    page.screenshot(path=str(output / f"{name}-failure.png"), full_page=True)
                    (output / f"{name}-failure.txt").write_text(page.locator("body").inner_text())
                    case["controls"] = page.locator("button,input,[role=combobox],[role=radio],summary").evaluate_all("""els=>els.map(e=>({
                        tag:e.tagName,role:e.getAttribute('role'),text:e.innerText?.slice(0,100),
                        label:e.getAttribute('aria-label'),type:e.getAttribute('type'),
                        visible:!!e.getClientRects().length}))""")
                finally:
                    case["console"] = r.console
                    case["failed_requests"] = r.requests
                    case["websockets"] = r.websockets
                    context.tracing.stop(path=str(output / f"{name}-trace.zip"))
                    context.close()
                    (output / "report.json").write_text(json.dumps(report, indent=2))
            browser.close()
    failures = sum(case.get("status") != "passed" for case in report["cases"])
    print(f"Browser acceptance: {len(report['cases']) - failures} passed, {failures} failed. Evidence: {output}", flush=True)
    if failures or len(report["cases"]) != 9:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
